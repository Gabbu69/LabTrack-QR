"""Build the two editable client documents. Use the Codex bundled Python runtime.

Hosted screenshots belong in output/playwright/client-guide. Render the resulting
DOCX files to PDF and inspect all pages before supplying either format.
"""
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "doc"
SHOTS = ROOT / "output" / "playwright" / "client-guide"
NAVY = "16374D"


def base(title, compact=False):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(15)
    section.left_margin = section.right_margin = Mm(17)
    section.footer_distance = Mm(8)
    styles = doc.styles
    # The supplied Word template can carry a blue rule in its Title style.
    for border in styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption", "List Bullet", "List Number"):
        styles[name].font.name = "Calibri"
        styles[name].font.color.rgb = RGBColor.from_string("172B38")
    normal = styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(5 if compact else 7)
    normal.paragraph_format.line_spacing = 1.03 if compact else 1.08
    styles["Title"].font.size = Pt(22 if compact else 25)
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    styles["Title"].paragraph_format.space_after = Pt(8)
    styles["Heading 1"].font.size = Pt(19)
    styles["Heading 1"].font.color.rgb = RGBColor.from_string(NAVY)
    styles["Heading 1"].paragraph_format.space_after = Pt(9)
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 2"].font.bold = True
    styles["Heading 2"].paragraph_format.space_before = Pt(9)
    styles["Heading 2"].paragraph_format.space_after = Pt(5)
    styles["Caption"].font.size = Pt(9)
    styles["Caption"].font.italic = False
    styles["Caption"].font.bold = False
    styles["Caption"].paragraph_format.space_after = Pt(9)
    for name in ("List Bullet", "List Number"):
        styles[name].font.size = Pt(11)
        styles[name].paragraph_format.space_after = Pt(5)
    doc.core_properties.title = title
    doc.core_properties.subject = "LabTrack QR client software and thesis demonstration handoff"
    doc.core_properties.author = "LabTrack QR"
    doc.core_properties.keywords = "client guide, thesis demo, payment proposal"
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("LabTrack QR  |  ")
    run.font.size = Pt(8.5)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    return doc


def para(doc, text, bold_lead=None, style=None):
    p = doc.add_paragraph(style=style)
    if bold_lead:
        p.add_run(bold_lead + " ").bold = True
    p.add_run(text)
    return p


def table(doc, headers, rows, widths, size=10.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for column, width in zip(t.columns, widths):
        column.width = Mm(width)
    for i, text in enumerate(headers):
        t.rows[0].cells[i].text = text
    for values in rows:
        cells = t.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = value
    for n, row in enumerate(t.rows):
        row_properties = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        row_properties.append(no_split)
        if n == 0:
            row_properties.append(OxmlElement("w:tblHeader"))
        for i, cell in enumerate(row.cells):
            cell.width = Mm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                border = OxmlElement("w:" + edge)
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    border.set(qn("w:" + key), value)
                borders.append(border)
            properties.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                margin = OxmlElement("w:" + edge)
                margin.set(qn("w:w"), "75")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            properties.append(margins)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), NAVY if n == 0 else ("F0F5F8" if n % 2 else "FFFFFF"))
            properties.append(shading)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.03
                for run in p.runs:
                    run.font.size = Pt(size)
                    if n == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def figure(doc, name, caption, width=166):
    path = SHOTS / name
    if not path.is_file():
        raise FileNotFoundError(f"Capture the real hosted screenshot first: {path}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Mm(width))
    para(doc, caption, style="Caption")


def page(doc, title):
    doc.add_page_break()
    doc.add_heading(title, level=1)


def build_guide():
    doc = base("LabTrack QR Client Guide")
    doc.add_paragraph("LabTrack QR Client Guide", "Title")
    para(doc, "A practical guide for one laboratory and three roles. Use it to rehearse QR-assisted tool custody and prepare the November thesis demonstration.")
    para(doc, "Website: https://labtrack-qr.vercel.app   User guide: /guide")
    figure(doc, "custodian-dashboard.png", "Hosted custodian demo on 9 October 2026. Names and inventory shown are fictional.")
    table(doc, ["Role", "What this role does"], [
        ("Custodian", "Approves accounts, manages tools, accepts checkout/return and exports reports."),
        ("Instructor", "Monitors inventory, borrowers and history; cannot perform custody or account changes."),
        ("Student", "Shows a personal QR and views only their own custody, history and profile."),
    ], [32, 144])
    doc.add_heading("Start with the demonstration", level=2)
    para(doc, "Open /guide, choose a role and select Use demo account. The sign-in page fills the shared demo details; press Sign in yourself. Check the yellow DEMO MODE banner. Demo sign-in addresses are LabTrack accounts, not Gmail inboxes. Coordinate shared-demo practice and do not reset someone else's rehearsal.")
    doc.add_heading("Personal account verification", level=2)
    para(doc, "The hosted registration page observed on 9 October uses the authenticator-default configuration. Follow its setup/current-code screen. Inbox codes apply only after the owner enables and verifies email-code login; existing authenticators stay required. This observation does not establish a real-phone or inbox-delivery test.")

    page(doc, "Prepare accounts and tool labels")
    figure(doc, "inventory.png", "Each inventory row represents one physical asset. Print labels and test the first label before producing a full set.")
    doc.add_heading("Approve the right student", level=2)
    para(doc, "Students enter their own details at /register and press Create account once. Follow any email-confirmation message and the verification screen shown. Registration can take the student directly to verification. The student remains in Approval pending until a custodian approves them for tool issue.")
    para(doc, "The custodian opens User Management, checks the name and Student ID against the intended student, and presses Approve / activate. Disable suspends access while preserving borrowing history. An already-used Student ID needs account recovery rather than another registration.")
    doc.add_heading("Prepare authorized staff", level=2)
    para(doc, "In User Management, open Create a staff account, choose the intended role and supply a one-time password of at least 10 characters. Give it privately. Personal staff finish configured verification and choose a private password before laboratory access. A password reset does not remove an enrolled authenticator.")
    doc.add_heading("Create and match physical assets", level=2)
    para(doc, "In Tool Inventory, open Add a tool batch. Enter the tool name, category, quantity, condition and prefix; press Create assets & labels. A quantity of five creates five separately identified physical tools. Open Print labels or Print A4 sheet, print at 100% scale, and attach each code to its matching asset. Check the actual phone, lighting and printed label before making more.")
    para(doc, "The QR identifies a record; it does not sign a person in. Keep passwords, sign-in codes and authenticator setup keys private.")

    page(doc, "Borrow and reconcile returns")
    figure(doc, "borrow-review.png", "Review the identified borrower and selected assets before confirming. This screenshot shows two selected tools; no checkout was submitted.")
    doc.add_heading("Checkout", level=2)
    para(doc, "Open Borrow Tools. Scan the student's QR or type their Student ID and choose Use code. Confirm the student, then scan or type each available physical asset code. Remove a mistaken selection, review the tools and press Confirm checkout once. Hand over tools only after the success message. Duplicate or unavailable tools must not be issued.")
    doc.add_heading("Partial and complete returns", level=2)
    para(doc, "Open Process Return and identify the same borrower. Scan only the tools physically received. Choose Good, Fair or Damaged and add a useful note; Keep unavailable prevents reissue. Confirm the received selection. Unscanned items remain outstanding. Identify the borrower again and receive the remaining tool to complete the return. Check custody/history after an interrupted connection before retrying.")
    doc.add_heading("Missing and recovered items", level=2)
    para(doc, "Open Explicitly mark missing only after a tool is confirmed missing. Select that item, enter a note, and confirm. A partial return alone never means missing. Already-missing items cannot be marked missing again. When found, scan and return the item normally; its missing incident stays in history. A damaged return should remain unavailable.")
    para(doc, "When a return has been accepted, do not repeat it. If the follow-up refresh fails, reload the borrower to check current custody before taking another action.")

    page(doc, "Rehearse the seven minute defense")
    figure(doc, "history.png", "History filters apply to Export CSV and Print A4 report. Reports include all matching records, with each tool's own return date.")
    table(doc, ["Time", "Demonstration"], [
        ("0:30", "Introduce one laboratory, three roles and fictional demo records."),
        ("0:45", "Show physical assets, an individual QR and the label sheet."),
        ("0:45", "Show student approval and the student's personal QR."),
        ("1:30", "Identify one student, review two physical tools and confirm checkout."),
        ("1:30", "Return one tool, explain partial custody, then accept the second."),
        ("1:00", "Mark a prepared borrowed tool missing, then show recovery."),
        ("0:45", "Switch to instructor monitoring, compare CSV/A4 and show the blank worksheet."),
    ], [20, 156], size=10.5)
    doc.add_heading("Prepare a fallback", level=2)
    para(doc, "Rehearse the actual phone and printed labels. Camera access needs HTTPS and permission. If lighting or permission fails, use a USB scanner, typed codes, or JPEG/PNG/WebP QR-image upload up to 5 MB. SVG downloads are printable artwork; use a supported image when uploading to the scanner.")
    para(doc, "For fictional practice, DEMO-2026-04 with DMM-002 and SDR-002 was available when reviewed. Recheck availability before rehearsal because the demo is shared. Use the existing blank worksheet at /thesis-evaluation.html; collect actual observations under the adviser's approved process. The system does not establish aviation airworthiness or a passing defense.")

    page(doc, "Acceptance troubleshooting and handoff")
    doc.add_heading("Resolve common problems", level=2)
    table(doc, ["Problem", "Action"], [
        ("Approval pending", "Ask the custodian to check/approve the student's correct account."),
        ("Email code absent", "First confirm email-code mode is enabled. Check spam; resend after 60 seconds. Codes expire after 10 minutes. Contact the owner for sender problems."),
        ("Authenticator unavailable", "Use a verified backup device or contact the owner for account-specific identity/recovery. A password reset does not remove the factor."),
        ("Stale custody or interruption", "Do not repeat an accepted return. Reload the borrower and check current custody/history before another action."),
        ("Camera denied or poor scan", "Use a supported QR image or the Student ID/asset-code input; test printed size and lighting."),
    ], [40, 136], size=10.5)
    doc.add_heading("Record acceptance when observed", level=2)
    for text in [
        "[ ] Three roles can sign in, navigate, refresh and log out; role restrictions hold.",
        "[ ] Two tools check out once; partial/full, damaged and missing/recovery scenarios are correct.",
        "[ ] Printed labels scan on the client's actual phone; image and typed-code fallbacks work.",
        "[ ] Filtered CSV and A4 reports agree; the evaluation worksheet remains blank until used.",
        "[ ] Client receives guides, training and source/setup after the balance is confirmed.",
    ]:
        para(doc, text)
    doc.add_heading("Name the owner and agree the dates", level=2)
    para(doc, "Proposed rehearsal: 13 November 2026. Proposed accepted handoff: 17 November. Intended use: around 20 November onward. One 60-minute training/rehearsal and one consolidated minor guide/presentation correction round are included. Bug fixes last 30 calendar days from actual accepted handoff; 17 November acceptance gives a proposed 17 December end date.")
    para(doc, "Hosting is currently managed through the developer's provider accounts. Record the responsible owner, alternate, agreed duration and separate provider costs. Do not share provider credentials. Keep source/configuration, database/Auth and private photo backups; CSV alone cannot restore the system. A real laboratory pilot needs a separate onboarding and disposable-target recovery trial.")
    para(doc, "Client and actual acceptance date: __________________________________")
    para(doc, "Owner and alternate: ______________________________________________")
    para(doc, "Source/deployment and private backup record: _________________________")
    para(doc, "Agreed hosting end date and support end date: ________________________")
    return doc


def build_proposal():
    doc = base("LabTrack QR Payment and Service Proposal", compact=True)
    doc.add_paragraph("LabTrack QR Payment and Service Proposal", "Title")
    para(doc, "Proposed software and handoff package for client acceptance. Total PHP 10,000. Intended thesis use around 20 November 2026 onward.")
    doc.add_heading("Included delivery", level=2)
    para(doc, "One laboratory and three roles; approval/verification, individual tool QR labels, checkout, partial/full/damaged returns, missing recovery, CSV/A4 reports and fictional demo access. Includes a five-page guide in PDF and editable Word, source/setup files after full payment, one 60-minute training/rehearsal and one consolidated round of minor guide/presentation corrections before handoff.")
    doc.add_heading("Milestone payments", level=2)
    table(doc, ["Milestone", "Amount", "Payment trigger"], [
        ("Confirm package", "PHP 4,000\n40%", "Client agrees to the written scope, schedule and service terms."),
        ("Successful demo", "PHP 4,000\n40%", "Client observes the agreed roles and core checkout/return/report workflows."),
        ("Accepted handoff", "PHP 2,000\n20%", "Agreed defects resolved; guides/setup/training ready. Release source when the balance is confirmed."),
        ("Total", "PHP 10,000", ""),
    ], [37, 29, 110], size=10.5)
    para(doc, "Pay through GCash or bank transfer; recipient details are supplied privately. After each actual payment, acknowledge the amount, date, reference and remaining balance in writing.")
    doc.add_heading("Schedule and support", level=2)
    para(doc, "Target rehearsal: 13 November 2026. Target accepted handoff: 17 November. Dates require client confirmation. Bug fixes for agreed features run for 30 calendar days from actual accepted handoff; 17 November acceptance gives a proposed end date of 17 December 2026.")
    para(doc, "New modules, major adviser-requested changes, manuscript work and continuing maintenance require a separate quote. Participant recruitment/results and defense approval remain the client's and adviser's responsibilities. Provider fees, if any, are separate and require agreement.")
    doc.add_heading("Acceptance and hosting", level=2)
    para(doc, "Accept after observing roles, two-tool checkout/partial/full return, damaged/missing recovery, report agreement and actual phone/printed-label/fallback checks. Record untested items as pending. Current hosting remains under the developer's accounts; agree its duration privately. Continuing hosting beyond defense/support needs a separate agreement. Real laboratory use needs separate pilot and recovery acceptance.")
    para(doc, "Client name/signature: ________________________  Date: __________")
    para(doc, "Developer name/signature: _____________________  Date: __________")
    para(doc, "Accepted handoff: __________  Support ends: __________")
    para(doc, "Hosting owner/alternate and agreed end date: _____________________")
    para(doc, "Payment acknowledgment   Received PHP ______ on ______; reference ______; balance PHP ______.")
    return doc


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for filename, builder in (("LabTrack-QR-Client-Guide.docx", build_guide), ("LabTrack-QR-Payment-and-Service-Proposal.docx", build_proposal)):
        path = OUT / filename
        builder().save(path)
        print(path)
